from io import BytesIO
from pathlib import Path

from screening_agent.ingestion.parsers.base_parser import BaseParser, ParsedDocument
from docx import Document


class DocxParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        document = Document(str(path))
        return self._build_parsed_document(document, path.name)

    def parse_bytes(self, filename: str, file_bytes: bytes) -> ParsedDocument:
        document = Document(BytesIO(file_bytes))
        return self._build_parsed_document(document, filename)

    def _build_parsed_document(self, document: Document, filename: str) -> ParsedDocument:
        raw_text = self.remove_duplicates(
            self.extract_text_from_docx(document)
        )
        return ParsedDocument(
            raw_text=raw_text,
            filename=filename,
        )

    @staticmethod
    def remove_duplicates(raw_text: str) -> str:
        lines = raw_text.split('\n')
        cleaned = []

        for line in lines:
            if not cleaned or line != cleaned[-1]:
                cleaned.append(line)
        return "\n".join(cleaned)

    @staticmethod
    def extract_text_from_docx(document: Document) -> str:

        parts: list[str] = []

        def normalize(text: str) -> str:
            return " ".join(text.split())

        for p in document.paragraphs:
            text = normalize(p.text)
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        text = normalize(p.text)
                        if text:
                            parts.append(text)

        return "\n".join(parts)






