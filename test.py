from pathlib import Path
from docx import Document


def remove_duplicates(raw_text: str) -> str:
    lines = raw_text.split('\n')
    cleaned = []

    for line in lines:
        if not cleaned or line != cleaned[-1]:
            cleaned.append(line)
    return "\n".join(cleaned)


def extract_text_from_docx(path: Path) -> str:
    document = Document(str(path))

    parts: list[str] = []

    def normalize(text: str) -> str:
        return " ".join(text.split())

    for p in document.paragraphs:
        text = normalize(p.text)
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = normalize(p.text)
                    if text:
                        parts.append(text)

    return "\n".join(parts)



path = Path("Асташкевич Задание.docx")
# print(remove_duplicates(extract_text_from_docx(Path("/Users/macbookairm1/Downloads/Асташкевич Задание.docx"))))
# print(remove_duplicates(extract_text_from_docx(Path("/Users/macbookairm1/Downloads/[Sr ML] Maryia_A_JAN_2026_2 (3).docx"))))
print(path.suffix)